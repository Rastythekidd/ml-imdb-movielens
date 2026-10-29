"""Odtworzenie wybranych wierszy z oryginalnych archiwów i zapisanego manifestu."""
import sys
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')

from pathlib import Path
import argparse, csv, hashlib, json, random, re, tarfile, urllib.request, zipfile
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
def sha(p):
    h=hashlib.sha256()
    with p.open('rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''):h.update(b)
    return h.hexdigest()
def read_xlsx(name):
    w=load_workbook(ROOT/'data/original'/name,read_only=True,data_only=True)
    rows=list(w.active.iter_rows(values_only=True));w.close();return rows[0],[list(r) for r in rows[1:]]
def save(p,header,rows):
    with p.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(header);w.writerows(rows)
def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--raw-dir',type=Path,default=ROOT/'data/raw')
    ap.add_argument('--download',action='store_true',help='Pobierz dwa oficjalne archiwa źródłowe.')
    ap.add_argument('--out-dir',type=Path,default=ROOT/'data/original')
    args=ap.parse_args();m=json.loads((ROOT/'sources/source_manifest.json').read_text(encoding='utf-8'))
    args.raw_dir.mkdir(parents=True,exist_ok=True);args.out_dir.mkdir(parents=True,exist_ok=True)
    for key,name in [('imdb','aclImdb_v1.tar.gz'),('movielens','ml-100k.zip')]:
        p=args.raw_dir/name
        if not p.exists() and args.download:
            req=urllib.request.Request(m['sources'][key]['archive_url'],headers={'User-Agent':'University-ML-Stage1/1.0'})
            with urllib.request.urlopen(req,timeout=120) as response,p.open('wb') as out:
                while b:=response.read(1024*1024):out.write(b)
        if not p.exists():raise FileNotFoundError(f'{p}: dodaj archiwum lub użyj --download.')
        assert sha(p)==m['sources'][key]['archive_sha256'],f'Niezgodna suma kontrolna archiwum: {name}'
    with zipfile.ZipFile(args.raw_dir/'ml-100k.zip') as z:
        allratings=[list(map(int,line.split('\t'))) for line in z.read('ml-100k/u.data').decode('ascii').splitlines()]
        films={int(f[0]):f for f in (line.split('|') for line in z.read('ml-100k/u.item').decode('latin1').splitlines())}
        (ROOT/'sources/MovieLens_README.txt').write_bytes(z.read('ml-100k/README'))
    assert random.Random(42).sample(range(len(allratings)),200)==m['movielens_source_indices']
    numeric=[];view=[]
    for i in m['movielens_source_indices']:
        uid,mid,rating,ts=allratings[i];year=int(films[mid][2][-4:])
        numeric.append([uid,mid,rating,ts,year]);view.append([mid,films[mid][1],rating,year,uid,ts])
    reviews=[]
    wanted={item['archive_file'] for item in m['imdb_rows']}
    wanted.update(['aclImdb/README','aclImdb/train/urls_pos.txt','aclImdb/train/urls_neg.txt'])
    contents={}
    with tarfile.open(args.raw_dir/'aclImdb_v1.tar.gz','r|gz') as t:
        for member in t:
            if member.name in wanted:
                contents[member.name]=t.extractfile(member).read()
            if len(contents)==len(wanted):break
    assert set(contents)==wanted
    (ROOT/'sources/IMDb_README.txt').write_bytes(contents['aclImdb/README'])
    urls={label:contents[f'aclImdb/train/urls_{label}.txt'].decode('utf-8').splitlines() for label in ['pos','neg']}
    for item in m['imdb_rows']:
        body=contents[item['archive_file']]
        assert hashlib.sha256(body).hexdigest()==item['text_sha256']
        label,rid,score=re.search(r'/(pos|neg)/(\d+)_(\d+)\.txt$',item['archive_file']).groups()
        sentiment='pozytywny' if label=='pos' else 'negatywny'
        assert sentiment==item['sentiment'] and int(score)==item['original_score']
        url=urls[label][int(rid)]
        assert url==item['original_review_url'] and re.search(r'tt\d+',url).group()==item['id_imdb']
        tid=item['id_imdb']
        reviews.append([item['title_snapshot'],body.decode('utf-8'),sentiment,tid,'https://www.imdb.com/title/'+tid+'/'])
    for name,rows in [('IMDb_recenzje.xlsx',reviews),('MovieLens_liczby.xlsx',numeric),('MovieLens_filmy.xlsx',view)]:
        canonical=json.dumps(rows,ensure_ascii=False,separators=(',',':')).encode('utf-8')
        assert hashlib.sha256(canonical).hexdigest()==m['canonical_rows_sha256'][name],f'Niezgodna suma kontrolna wierszy: {name}'
        if (ROOT/'data/original'/name).exists():
            header,expected=read_xlsx(name)
            assert rows==expected,f'Odtworzone wartości różnią się: {name}'
        else:
            header={'IMDb_recenzje.xlsx':['tytul_filmu','tekst_recenzji','sentyment','id_imdb','link_imdb'],
                'MovieLens_liczby.xlsx':['id_uzytkownika','id_filmu','ocena','czas_unix','rok_premiery'],
                'MovieLens_filmy.xlsx':['id_filmu','tytul_filmu','ocena','rok_premiery','id_uzytkownika','czas_unix']}[name]
        save(args.out_dir/(Path(name).stem+'.csv'),header,rows)
    (args.out_dir/'verification.json').write_text(json.dumps({'review_rows':200,'rating_rows':200,'all_values_match_manifest':True,'source_checksums_verified':True},indent=2),encoding='utf-8')
    print('Potwierdzono archiwa i wartości: 200 recenzji oraz 200 ocen; porównano również dostępne arkusze Excel.')
    print('Nazwy IMDb odtworzono z zapisanego snapshotu metadanych w manifeście.')
if __name__=='__main__':main()
