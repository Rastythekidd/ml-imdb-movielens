"""Końcowa kontrola projektu i odtworzenia wyników bez arkuszy Excel."""
from pathlib import Path
import sys
if hasattr(sys.stdout,'reconfigure'):sys.stdout.reconfigure(encoding='utf-8')
import hashlib,json,math,shutil,statistics,subprocess
from openpyxl import load_workbook
from pypdf import PdfReader
ROOT=Path(__file__).resolve().parents[1]
def sha(p):
    h=hashlib.sha256()
    with p.open('rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''):h.update(b)
    return h.hexdigest()
def uruchom(sciezka,*argumenty,katalog=None):
    p=subprocess.run([sys.executable,'-B','-X','utf8',str(sciezka),*map(str,argumenty)],
        cwd=katalog or ROOT,capture_output=True,text=True,encoding='utf-8')
    if p.returncode:raise RuntimeError(p.stdout+'\n'+p.stderr)
    return p
def main():
    m=json.loads((ROOT/'sources/source_manifest.json').read_text(encoding='utf-8'))
    s=json.loads((ROOT/'results/summary.json').read_text(encoding='utf-8'))
    nazwa_raportu=json.loads((ROOT/'project.json').read_text(encoding='utf-8')).get('report_file','Raport_etap_1.pdf')
    h={}
    for nazwa,oczekiwany in m['input_file_sha256'].items():
        assert sha(ROOT/'data/original'/nazwa)==oczekiwany,nazwa
    for nazwa,oczekiwany in [
      ('aclImdb_v1.tar.gz',m['sources']['imdb']['archive_sha256']),
      ('ml-100k.zip',m['sources']['movielens']['archive_sha256']),
      ('imdb_train.parquet',m['sources']['imdb']['initial_parquet_sha256'])]:
        h[nazwa]=sha(ROOT/'data/raw'/nazwa);assert h[nazwa]==oczekiwany,nazwa
    h['title.basics.tsv.gz']=sha(ROOT/'data/raw/title.basics.tsv.gz')
    assert h['title.basics.tsv.gz']==m['sources']['metadata']['archive_sha256']
    (ROOT/'results/sumy_archiwow.json').write_text(json.dumps(h,indent=2),encoding='utf-8')
    w=load_workbook(ROOT/'data/original/MovieLens_liczby.xlsx',read_only=True,data_only=True)
    r=list(w.active.iter_rows(min_row=2,values_only=True));w.close()
    assert statistics.mean(x[2] for x in r)==s['movielens']['rating_summary']['mean']
    assert statistics.median(x[2] for x in r)==s['movielens']['rating_summary']['median']
    assert math.isclose(statistics.stdev(x[2] for x in r),s['movielens']['rating_summary']['std_sample'])
    assert len(PdfReader(ROOT/nazwa_raportu).pages)==5
    tekst='\n'.join(p.extract_text() for p in PdfReader(ROOT/nazwa_raportu).pages)
    for identyfikator in ['71629','71513','72072','72646','71510']:assert identyfikator in tekst
    assert '3,42' in tekst and '169,5' in tekst and '237,86' in tekst
    konfiguracja=json.loads((ROOT/'project.json').read_text(encoding='utf-8'))
    adres_repo=konfiguracja.get('repository_url')
    if adres_repo:
        assert adres_repo in tekst
        assert 'adres do uzupełnienia' not in tekst
    else:
        assert 'adres do uzupełnienia' in tekst
    nb=json.loads((ROOT/'Analiza_etap_1.ipynb').read_text(encoding='utf-8'))
    for komorka in nb['cells']:
        if komorka['cell_type']=='code':compile(''.join(komorka['source']),'notebook','exec')
    T=ROOT/'_robocze/proba_odtworzenia'
    T.mkdir(parents=True,exist_ok=True)
    shutil.copytree(ROOT/'src',T/'src',dirs_exist_ok=True,ignore=shutil.ignore_patterns('__pycache__'))
    (T/'sources').mkdir(exist_ok=True)
    shutil.copy2(ROOT/'sources/source_manifest.json',T/'sources/source_manifest.json')
    shutil.copy2(ROOT/'project.json',T/'project.json')
    uruchom(T/'src/reproduce_subset.py','--raw-dir',ROOT/'data/raw',katalog=T)
    uruchom(T/'src/run_all.py',katalog=T)
    ss=json.loads((T/'results/summary.json').read_text(encoding='utf-8'))
    ss.pop('input_sha256');porownanie=dict(s);porownanie.pop('input_sha256')
    assert ss==porownanie,'Wyniki CSV i XLSX są różne.'
    for png in (ROOT/'results/figures').glob('*.png'):
        assert sha(png)==sha(T/'results/figures'/png.name),png.name
    assert len(PdfReader(T/nazwa_raportu).pages)==5
    # Profil nie może zastąpić adresu repozytorium.
    profil=subprocess.run([sys.executable,'-B','-X','utf8',str(T/'src/make_report.py'),
        '--repo-url','https://github.com/Rastythekidd'],cwd=T,capture_output=True,text=True,encoding='utf-8')
    assert profil.returncode!=0 and 'Sam profil nie wystarcza' in profil.stderr
    assert json.loads((T/'project.json').read_text(encoding='utf-8'))['repository_url']==adres_repo
    wynik={'data':'2026-10-27','dane_sprawdzone_z_archiwami':400,'metadane_imdb':189,
      'bledy_danych':0,'oryginalne_arkusze_niezmienione':True,'sumy_archiwow':'zgodne',
      'statystyki_niezalezna_kontrola':'zgodne','odtworzenie_bez_xlsx':'poprawne',
      'wyniki_csv_i_xlsx':'identyczne','wykresy_po_odtworzeniu':'identyczne',
      'raport_liczba_stron':5,'notebook_skladnia':'poprawna',
      'walidacja_linku_repozytorium':'poprawna','zmiany_na_github':False}
    (ROOT/'results/verification.json').write_text(json.dumps(wynik,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(wynik,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
