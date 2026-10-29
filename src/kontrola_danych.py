"""Porównanie wszystkich rekordów z oryginalnymi archiwami i metadanymi."""
import sys
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')

from pathlib import Path
from collections import Counter
import hashlib, json, tarfile, zipfile, gzip, re, statistics
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
P=ROOT
SRC=ROOT/'data/original'
RAW=ROOT/'data/raw'
m=json.loads((P/'sources/source_manifest.json').read_text(encoding='utf-8'))
sha=lambda b:hashlib.sha256(b).hexdigest()
books={};checks={};issues=[]
for name in m['input_file_sha256']:
 w=load_workbook(SRC/name,read_only=True,data_only=False)
 rows=list(w.active.iter_rows(values_only=True));w.close();books[name]=rows
 checks[name]={'wiersze':len(rows)-1,'kolumny':len(rows[0]),'braki':sum(v is None or (isinstance(v,str) and not v.strip()) for row in rows[1:] for v in row),'duplikaty':len(rows[1:])-len(set(rows[1:])),'formuly':sum(isinstance(v,str) and v.startswith('=') for row in rows[1:] for v in row),'zgodny_hash':sha((SRC/name).read_bytes())==m['input_file_sha256'][name]}
 if not checks[name]['zgodny_hash']:issues.append({'plik':name,'problem':'Plik zmieniony od poprzedniego audytu'})
with zipfile.ZipFile(RAW/'ml-100k.zip') as z:
 ratings=[list(map(int,s.split('\t'))) for s in z.read('ml-100k/u.data').decode('ascii').splitlines()]
 films={int(f[0]):f for f in (s.split('|') for s in z.read('ml-100k/u.item').decode('latin1').splitlines())}
expected_num=[];expected_view=[]
for i in m['movielens_source_indices']:
 uid,mid,rating,ts=ratings[i];year=int(films[mid][2][-4:])
 expected_num.append([uid,mid,rating,ts,year]);expected_view.append([mid,films[mid][1],rating,year,uid,ts])
for name,expected in [('MovieLens_liczby.xlsx',expected_num),('MovieLens_filmy.xlsx',expected_view)]:
 for i,(actual,wanted) in enumerate(zip(books[name][1:],expected)):
  for j,(a,b) in enumerate(zip(actual,wanted)):
   if a!=b:issues.append({'plik':name,'wiersz':i+2,'kolumna':j+1,'jest':a,'powinno':b})
wanted={r['archive_file'] for r in m['imdb_rows']}|{'aclImdb/train/urls_pos.txt','aclImdb/train/urls_neg.txt'}
contents={}
with tarfile.open(RAW/'aclImdb_v1.tar.gz','r|gz') as t:
 for member in t:
  if member.name in wanted:contents[member.name]=t.extractfile(member).read()
  if len(contents)==len(wanted):break
urls={c:contents[f'aclImdb/train/urls_{c}.txt'].decode().splitlines() for c in ['pos','neg']}
for i,item in enumerate(m['imdb_rows']):
 body=contents[item['archive_file']]
 label,rid,score=re.search(r'/(pos|neg)/(\d+)_(\d+)\.txt$',item['archive_file']).groups()
 sentiment='pozytywny' if label=='pos' else 'negatywny'
 actual=list(books['IMDb_recenzje.xlsx'][i+1])
 expected=[item['title_snapshot'],body.decode(),sentiment,item['id_imdb'],'https://www.imdb.com/title/'+item['id_imdb']+'/']
 if actual!=expected:issues.append({'plik':'IMDb_recenzje.xlsx','wiersz':i+2,'problem':'Niezgodnosc ze zrodlem'})
 if sha(body)!=item['text_sha256']:issues.append({'wiersz':i+2,'problem':'Hash tekstu'})
 if re.search(r'tt\d+',urls[label][int(rid)]).group()!=item['id_imdb']:issues.append({'wiersz':i+2,'problem':'ID tytulu'})
expected_meta={r['id_imdb']:r['title_metadata'] for r in m['imdb_rows']};seen=set()
with gzip.open(RAW/'title.basics.tsv.gz','rt',encoding='utf-8') as f:
 next(f)
 for line in f:
  tid=line.split('\t',1)[0]
  if tid not in expected_meta:continue
  fields=line.rstrip('\n').split('\t');e=expected_meta[tid]
  if [fields[2],fields[3],fields[5],fields[1]]!=[e['title'],e['original_title'],e['year'],e['title_type']]:issues.append({'id':tid,'problem':'Metadane'})
  seen.add(tid)
  if len(seen)==len(expected_meta):break
result={'kontrola':checks,'liczba_problemow_danych':len(issues),'problemy':issues,'zweryfikowane_metadane':len(seen),'srednia_ocena':statistics.mean(r[2] for r in expected_num),'mediana_oceny':statistics.median(r[2] for r in expected_num),'klasy_imdb':dict(Counter(row[2] for row in books['IMDb_recenzje.xlsx'][1:]))}
(ROOT/'results/kontrola_danych_szczegolowa.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
if issues: raise SystemExit('Wykryto niezgodności w danych.')
